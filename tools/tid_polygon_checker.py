r"""Standalone local Windows TID polygon boundary checker.

Input is UTF-8 text containing one ``TrackingID<TAB>Station`` pair per line.
Every non-empty input line produces exactly one output row in the same order;
duplicates and the same TID paired with different stations are preserved.

The tool is read-only against SIW. It probes mcscli first and automatically
switches once to same-origin fetch in the authenticated local Edge profile when
mcscli is unavailable. No AgentSpace or local HTTP server is used.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable
from uuid import uuid4

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import gcrs_mcscli
import polygon_check
from siw_client import FailoverSiwSession, TransportError

RESULT_HEADERS = ["Tracking ID", "Station", "GCRS Status"]
ERROR_PREFIX = "Error — "
CAPTURE_DIR = Path(__file__).parent / "captures" / "polygon_local"
CAPTURE_DIR.mkdir(parents=True, exist_ok=True)


@dataclass(frozen=True)
class BoundaryInput:
    sequence: int
    source_line: int
    raw_line: str
    tracking_id: str
    station: str
    canonical_tid: str
    canonical_station: str
    carrier: str | None
    parse_error: str | None = None


@dataclass(frozen=True)
class BoundaryResult:
    source: BoundaryInput
    status: str

    def as_row(self) -> dict[str, str]:
        return {
            "Tracking ID": self.source.tracking_id,
            "Station": self.source.station,
            "GCRS Status": self.status,
        }


@dataclass
class BoundaryRun:
    inputs: list[BoundaryInput]
    results: list[BoundaryResult]
    gcrs_rows: list[dict]
    not_found: list[str]
    errored: list[str]
    summary: dict[str, object]
    transport: dict[str, object]


def _canonical_tid(value: str) -> tuple[str, str | None]:
    tid = value.strip().lstrip("\ufeff")
    if re.fullmatch(r"\d{12}", tid):
        return tid, "ATS"
    if re.fullmatch(r"DOC\d+", tid, re.IGNORECASE):
        return tid.upper(), "ATA"
    return tid, None


def parse_input_text(text: str) -> list[BoundaryInput]:
    inputs: list[BoundaryInput] = []
    for source_line, raw_line in enumerate(text.splitlines(), start=1):
        raw_line = raw_line.rstrip("\r")
        if not raw_line.strip():
            continue
        fields = raw_line.split("\t")
        parse_error = None
        if len(fields) != 2:
            tracking_id = raw_line.strip()
            station = ""
            parse_error = "expected exactly one tab between TID and station"
        else:
            tracking_id, station = (field.strip() for field in fields)
        canonical_tid, carrier = _canonical_tid(tracking_id)
        if parse_error is None and carrier is None:
            parse_error = "TID must be 12 digits or DOC followed by digits"
        inputs.append(
            BoundaryInput(
                sequence=len(inputs),
                source_line=source_line,
                raw_line=raw_line,
                tracking_id=tracking_id,
                station=station,
                canonical_tid=canonical_tid,
                canonical_station=station.strip().upper(),
                carrier=carrier,
                parse_error=parse_error,
            )
        )
    return inputs


def build_inputs_from_pairs(pairs: list[tuple[object, object]]) -> list[BoundaryInput]:
    return parse_input_text("\n".join(f"{'' if tid is None else tid}\t{'' if station is None else station}" for tid, station in pairs))


def run_boundary_check(
    inputs: list[BoundaryInput],
    session: FailoverSiwSession,
    log: Callable[[str], None] = print,
    refresh_polygons: bool = False,
    gcrs_result: tuple[list[dict], list[str], list[str]] | None = None,
) -> BoundaryRun:
    valid_tids = [item.canonical_tid for item in inputs if item.carrier and not item.parse_error]
    transport_error = ""
    if gcrs_result is None:
        try:
            gcrs_rows, not_found, errored = gcrs_mcscli.fetch(valid_tids, log=log, session=session)
        except (TransportError, RuntimeError, ValueError, TypeError) as exc:
            gcrs_rows, not_found, errored = [], [], list(dict.fromkeys(valid_tids))
            transport_error = str(exc)[:240]
    else:
        gcrs_rows, not_found, errored = gcrs_result

    gcrs_by_tid = {gcrs_mcscli.normalize_tid(row.get("Query_Item")): row for row in gcrs_rows}
    not_found_set = {gcrs_mcscli.normalize_tid(value) for value in not_found}
    errored_set = {gcrs_mcscli.normalize_tid(value) for value in errored}

    eligible: list[BoundaryInput] = []
    for item in inputs:
        if (
            item.parse_error is None
            and item.canonical_station
            and item.canonical_tid in gcrs_by_tid
            and item.canonical_tid not in not_found_set
            and item.canonical_tid not in errored_set
        ):
            eligible.append(item)
    polygon_rows = polygon_check.check_pairs(
        [(item.canonical_tid, item.canonical_station) for item in eligible],
        gcrs_rows,
        log=log,
        force_refresh=refresh_polygons,
        session=session,
    ) if eligible else []
    polygon_status = {item.sequence: row["GCRS Status"] for item, row in zip(eligible, polygon_rows)}

    results: list[BoundaryResult] = []
    for item in inputs:
        if item.parse_error:
            status = f"{ERROR_PREFIX}Invalid input line: {item.parse_error}"
        elif not item.canonical_station:
            status = f"{ERROR_PREFIX}Missing station code"
        elif transport_error:
            status = f"{ERROR_PREFIX}SIW unavailable: {transport_error}"
        elif item.canonical_tid in not_found_set:
            status = f"{ERROR_PREFIX}TID not found in SIW"
        elif item.canonical_tid in errored_set or item.canonical_tid not in gcrs_by_tid:
            status = f"{ERROR_PREFIX}SIW unresolved after degraded/poison guards"
        else:
            status = polygon_status.get(item.sequence, f"{ERROR_PREFIX}Polygon result missing")
        results.append(BoundaryResult(source=item, status=status))

    inside = sum(result.status == polygon_check.INSIDE_TEXT for result in results)
    outside = sum(result.status == polygon_check.OUTSIDE_TEXT for result in results)
    errors = sum(result.status.startswith(ERROR_PREFIX) for result in results)
    summary = {
        "Total": len(results),
        "Inside": inside,
        "Outside": outside,
        "Errors": errors,
        "Unique lookup TIDs": len(dict.fromkeys(valid_tids)),
        "Unique carrier/station pairs": len({(item.carrier, item.canonical_station) for item in inputs if item.carrier and item.canonical_station}),
    }
    return BoundaryRun(
        inputs=inputs,
        results=results,
        gcrs_rows=gcrs_rows,
        not_found=not_found,
        errored=errored,
        summary=summary,
        transport=session.summary(),
    )


def _set_text(cell, value: object) -> None:
    cell.value = "" if value is None else str(value)
    cell.data_type = "s"


def _default_output(run: BoundaryRun) -> Path:
    stations = len({item.canonical_station for item in run.inputs if item.canonical_station})
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return Path.home() / "Desktop" / f"TID_Check_BULK_{stations}stns_{stamp}.xlsx"


def _status_fill(status: str) -> PatternFill:
    if status == polygon_check.INSIDE_TEXT:
        return PatternFill("solid", fgColor="C6EFCE")
    if status == polygon_check.OUTSIDE_TEXT:
        return PatternFill("solid", fgColor="FFC7CE")
    return PatternFill("solid", fgColor="FFEB9C")


def write_results_xlsx(
    run: BoundaryRun,
    output_path: Path | None = None,
    overwrite: bool = False,
) -> tuple[Path, dict[str, object]]:
    output = Path(output_path) if output_path else _default_output(run)
    if len(run.results) != len(run.inputs):
        raise RuntimeError(
            f"refusing to write incomplete result: {len(run.results)} rows for {len(run.inputs)} inputs"
        )
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists() and not overwrite:
        raise FileExistsError(f"Output already exists: {output}")

    workbook = openpyxl.Workbook()
    results_sheet = workbook.active
    results_sheet.title = "Results"
    results_sheet.append(RESULT_HEADERS)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in results_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    for result in run.results:
        row_number = results_sheet.max_row + 1
        row = result.as_row()
        for column, header in enumerate(RESULT_HEADERS, start=1):
            _set_text(results_sheet.cell(row=row_number, column=column), row[header])
        results_sheet.cell(row=row_number, column=3).fill = _status_fill(result.status)
    results_sheet.freeze_panes = "A2"
    results_sheet.auto_filter.ref = f"A1:C{max(1, results_sheet.max_row)}"
    for index, width in enumerate((24, 16, 72), start=1):
        results_sheet.column_dimensions[get_column_letter(index)].width = width

    summary_sheet = workbook.create_sheet("Summary")
    summary_sheet.append(["Metric", "Value"])
    summary_rows = dict(run.summary)
    summary_rows.update({
        "Selected transport": run.transport.get("selected_transport") or "not required",
        "Fallback count": run.transport.get("fallback_count", 0),
        "Fallback reason": run.transport.get("fallback_reason", ""),
        "Rows == Input": f"{len(run.results)} == {len(run.inputs)}",
    })
    for key, value in summary_rows.items():
        summary_sheet.append([key, value])
    for cell in summary_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 80

    temporary = output.with_name(f".{output.stem}.{uuid4().hex}.tmp.xlsx")
    try:
        workbook.save(temporary)
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    verification = verify_results_xlsx(output, run)
    return output, verification


def verify_results_xlsx(path: Path, expected: BoundaryRun) -> dict[str, object]:
    if len(expected.results) != len(expected.inputs):
        raise RuntimeError(
            f"VERIFY FAIL: in-memory result/input mismatch {len(expected.results)} != {len(expected.inputs)}"
        )
    for index, (source, result) in enumerate(zip(expected.inputs, expected.results)):
        if result.source != source or source.sequence != index:
            raise RuntimeError(f"VERIFY FAIL: in-memory sequence mismatch at row {index + 1}")

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if "Results" not in workbook.sheetnames or "Summary" not in workbook.sheetnames:
            raise RuntimeError("VERIFY FAIL: Results or Summary sheet missing")
        results_sheet = workbook["Results"]
        headers = [cell.value for cell in next(results_sheet.iter_rows(min_row=1, max_row=1))]
        if headers != RESULT_HEADERS:
            raise RuntimeError(f"VERIFY FAIL: headers differ: {headers}")
        actual_rows = [
            tuple("" if value is None else str(value) for value in row)
            for row in results_sheet.iter_rows(min_row=2, values_only=True)
            if any(value is not None for value in row)
        ]
        expected_rows = [
            (result.source.tracking_id, result.source.station, result.status)
            for result in expected.results
        ]
        if len(actual_rows) != len(expected.inputs):
            raise RuntimeError(
                f"VERIFY FAIL: Excel/input row mismatch {len(actual_rows)} != {len(expected.inputs)}"
            )
        if actual_rows != expected_rows:
            mismatch = next(
                (index for index, pair in enumerate(zip(actual_rows, expected_rows), start=1) if pair[0] != pair[1]),
                None,
            )
            raise RuntimeError(
                f"VERIFY FAIL: ordered rows differ; actual={len(actual_rows)} "
                f"expected={len(expected_rows)} first={mismatch}"
            )

        statuses = [row[2] for row in actual_rows]
        valid_status = lambda status: (
            status in {polygon_check.INSIDE_TEXT, polygon_check.OUTSIDE_TEXT}
            or status.startswith(ERROR_PREFIX)
        )
        if not all(valid_status(status) for status in statuses):
            raise RuntimeError("VERIFY FAIL: unexpected status text")
        recomputed = {
            "Total": len(statuses),
            "Inside": sum(status == polygon_check.INSIDE_TEXT for status in statuses),
            "Outside": sum(status == polygon_check.OUTSIDE_TEXT for status in statuses),
            "Errors": sum(status.startswith(ERROR_PREFIX) for status in statuses),
        }
        if recomputed["Inside"] + recomputed["Outside"] + recomputed["Errors"] != recomputed["Total"]:
            raise RuntimeError("VERIFY FAIL: status counts do not cover every result row")

        summary_sheet = workbook["Summary"]
        summary = {
            str(key): value
            for key, value in summary_sheet.iter_rows(min_row=2, min_col=1, max_col=2, values_only=True)
            if key is not None
        }
        for key, value in recomputed.items():
            if int(summary.get(key, -1)) != value or int(expected.summary.get(key, -1)) != value:
                raise RuntimeError(f"VERIFY FAIL: independently recomputed summary mismatch for {key}")
        expected_row_statement = f"{len(actual_rows)} == {len(expected.inputs)}"
        if summary.get("Rows == Input") != expected_row_statement:
            raise RuntimeError("VERIFY FAIL: Rows == Input summary is incorrect")
    finally:
        workbook.close()

    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    return {
        "verify_ok": True,
        "rows": len(actual_rows),
        "input_rows": len(expected.inputs),
        "sha256": digest,
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Local Windows GCRS polygon boundary checker")
    parser.add_argument("input", help="UTF-8 file containing TID<TAB>Station pairs")
    parser.add_argument("--output", help="XLSX path; defaults to a timestamped file on Desktop")
    parser.add_argument("--refresh-polygons", action="store_true", help="ignore cached station polygons")
    parser.add_argument("--transport", choices=("auto", "mcscli", "edge"), default="auto")
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()

    input_path = Path(args.input).resolve()
    text = input_path.read_text(encoding="utf-8-sig")
    inputs = parse_input_text(text)
    if not inputs:
        raise RuntimeError("Input contains no non-empty lines")
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    log_path = CAPTURE_DIR / f"polygon_run_{stamp}.log"

    def log(message: str) -> None:
        line = f"[{datetime.now():%H:%M:%S}] {message}"
        print(line, flush=True)
        with open(log_path, "a", encoding="utf-8") as file:
            file.write(line + "\n")

    log(f"Parsed {len(inputs)} non-empty input rows")
    with FailoverSiwSession(mode=args.transport, log=log) as session:
        run = run_boundary_check(inputs, session, log=log, refresh_polygons=args.refresh_polygons)
    output, verification = write_results_xlsx(
        run,
        Path(args.output).resolve() if args.output else None,
        overwrite=args.overwrite,
    )
    log(f"SAVED {output}")
    log(
        f"Total={run.summary['Total']} Inside={run.summary['Inside']} Outside={run.summary['Outside']} "
        f"Errors={run.summary['Errors']}"
    )
    log(f"Rows == Input: {verification['rows']} == {verification['input_rows']} (VERIFY OK)")
    print(json.dumps({"output": str(output), "summary": run.summary, "transport": run.transport, "verification": verification}, indent=2))
    return 0


if __name__ == "__main__":
    sys.exit(main())
