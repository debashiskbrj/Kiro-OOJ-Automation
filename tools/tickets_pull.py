r"""Tech_Data producer: pull all open SWA OOJ tickets from SIM-T and explode to one TID per row.

Local port of the user's AgentSpace "SWA OOJ Ticket Analyzer" (same CTI filter, same
11-column schema, same extraction rules). Data path = builder-mcp.exe spawned as a stdio
MCP subprocess -> TicketingReadActions / search-tickets. Read-only.

Output (tools\captures\tickets\):
  Tech_Data_<YYYYMMDD_HHMM>.xlsx  (11 cols: Sl, case_id, createdAt, lastUpdatedAt,
                                   station_code, zone, impacted_tracking_ids, attachments,
                                   status, IssueUrl, ShortId)
  Tech_Data_<YYYYMMDD_HHMM>.csv   (same rows, for the pipeline)
  tickets_pull.log

Usage:  python tickets_pull.py [--max-pages N] [--keep-old]
Requires: a valid Midway session (mwinit -o) for builder-mcp.
"""
import argparse
import csv
import json
import queue
import re
import subprocess
import sys
import threading
import time
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

BUILDER_MCP = r"C:\Users\kabideba\AppData\Local\Toolbox\bin\builder-mcp.exe"
OUT = Path(__file__).parent / "captures" / "tickets"
OUT.mkdir(parents=True, exist_ok=True)
LOG = OUT / "tickets_pull.log"

# --- CTI filter (authoritative, matches the user's saved t.corp query) ---------------
ASSIGNED_GROUPS = ["SWA India", "SWA-CO-NOC", "in-swa-co"]
STATUSES = ["Assigned", "Researching", "Work In Progress", "Pending"]
QUERY = ('extensions.tt.category:"SWA" AND extensions.tt.type:"Out of Jurisdiction" '
         'AND extensions.tt.item:"Pin code not serviceable by station"')
RESPONSE_FIELDS = ["id", "aliases", "title", "status", "createDate", "lastUpdatedDate",
                   "extensions.tt.status", "extensions.tt.assignedGroup", "description", "attachments"]
PAGE_SIZE = 100

# --- extraction rules (exact, from the AgentSpace context.md) -----------------------
RE_TID = re.compile(r"DOC\d+|\b\d{12}\b")
RE_ZONE = re.compile(r"Zone:\s*(\w+)")
RE_STATION = re.compile(r"Station\s*code\s*:\s*([A-Za-z0-9]+)", re.IGNORECASE)
RE_ALIAS = re.compile(r"^[A-Z]\d+$")

COLUMNS = ["Sl", "case_id", "createdAt", "lastUpdatedAt", "station_code", "zone",
           "impacted_tracking_ids", "attachments", "status", "IssueUrl", "ShortId"]


def log(msg: str):
    line = f"[{datetime.now():%H:%M:%S}] {msg}"
    with open(LOG, "a", encoding="utf-8") as fh:
        fh.write(line + "\n")
    print(line, flush=True)


# --- minimal stdio MCP client ----------------------------------------------------------
class McpStdio:
    def __init__(self, exe: str, args: list[str]):
        self.proc = subprocess.Popen([exe, *args], stdin=subprocess.PIPE, stdout=subprocess.PIPE,
                                     stderr=subprocess.PIPE, text=True, encoding="utf-8", bufsize=1)
        self._id = 0
        self._stderr_lines = []
        self._stdout_queue = queue.Queue()
        threading.Thread(target=self._drain_stderr, daemon=True).start()
        threading.Thread(target=self._drain_stdout, daemon=True).start()

    def _drain_stderr(self):
        for line in self.proc.stderr:
            self._stderr_lines.append(line.rstrip())

    def _drain_stdout(self):
        for line in self.proc.stdout:
            self._stdout_queue.put(line)
        self._stdout_queue.put(None)

    def _send(self, obj: dict):
        self.proc.stdin.write(json.dumps(obj) + "\n")
        self.proc.stdin.flush()

    def request(self, method: str, params: dict | None = None, timeout: float = 300) -> dict:
        self._id += 1
        rid = self._id
        self._send({"jsonrpc": "2.0", "id": rid, "method": method, "params": params or {}})
        deadline = time.time() + timeout
        while True:
            remaining = deadline - time.time()
            if remaining <= 0:
                raise TimeoutError(f"{method} timed out after {timeout}s")
            try:
                line = self._stdout_queue.get(timeout=remaining)
            except queue.Empty as exc:
                raise TimeoutError(f"{method} timed out after {timeout}s") from exc
            if line is None:
                raise RuntimeError(f"builder-mcp closed stdout. stderr tail: {self._stderr_lines[-5:]}")
            line = line.strip()
            if not line.startswith("{"):
                continue
            try:
                msg = json.loads(line)
            except json.JSONDecodeError:
                continue
            if msg.get("id") == rid:
                if "error" in msg:
                    raise RuntimeError(f"{method} error: {msg['error']}")
                return msg.get("result", {})

    def notify(self, method: str, params: dict | None = None):
        self._send({"jsonrpc": "2.0", "method": method, "params": params or {}})

    def call_tool(self, name: str, arguments: dict, timeout: float = 300) -> dict:
        res = self.request("tools/call", {"name": name, "arguments": arguments}, timeout=timeout)
        content = res.get("content") or []
        text = "".join(c.get("text", "") for c in content if c.get("type") == "text")
        if res.get("isError"):
            raise RuntimeError(f"tool {name} returned error: {text[:500]}")
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            raise RuntimeError(f"tool {name} returned non-JSON: {text[:500]}")

    def close(self):
        try:
            if self.proc.stdin and not self.proc.stdin.closed:
                self.proc.stdin.close()
            self.proc.terminate()
            self.proc.wait(timeout=10)
        except subprocess.TimeoutExpired:
            self.proc.kill()
            self.proc.wait(timeout=5)
        except Exception:
            pass


def open_mcp() -> McpStdio:
    mcp = McpStdio(BUILDER_MCP, ["--include-tools=TicketingReadActions"])
    init = mcp.request("initialize", {"protocolVersion": "2025-03-26", "capabilities": {},
                                      "clientInfo": {"name": "ooj-tickets-pull", "version": "0.1"}}, timeout=120)
    mcp.notify("notifications/initialized")
    log(f"builder-mcp ready: {init.get('serverInfo')}")
    return mcp


# --- fetch --------------------------------------------------------------------------------
def fetch_all(mcp: McpStdio, max_pages: int | None) -> tuple[list[dict], int, int]:
    snapshot = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    query = f"{QUERY} AND createDate:[* TO {snapshot}]"
    log(f"snapshot upper bound: {snapshot}")

    tickets: dict[str, dict] = {}
    token = None
    total = None
    page = 0
    while True:
        page += 1
        inp = {"assignedGroup": ASSIGNED_GROUPS, "status": STATUSES, "query": query,
               "responseFields": RESPONSE_FIELDS, "sort": "createDate desc", "rows": PAGE_SIZE}
        if token:
            inp["startToken"] = token
        data = mcp.call_tool("TicketingReadActions", {"action": "search-tickets", "input": inp})
        if data.get("status") != "success":
            raise RuntimeError(f"search-tickets failed: {json.dumps(data)[:500]}")
        d = data["data"]
        total = d.get("totalResults", total)
        batch = d.get("tickets") or []
        new_token = d.get("startToken")
        for t in batch:
            tickets[t["id"]] = t
        log(f"page {page}: got {len(batch)}, collected {len(tickets)} / totalResults {total}")
        if not batch or (total is not None and len(tickets) >= total):
            break
        if not new_token or new_token == token:
            log("startToken did not advance — stopping pagination")
            break
        token = new_token
        if max_pages and page >= max_pages:
            log(f"stopping at --max-pages {max_pages}")
            break
    ordered = sorted(tickets.values(), key=lambda t: t.get("createDate", ""), reverse=True)
    return ordered, total or 0, page


# --- parse / explode -----------------------------------------------------------------------
def parse_iso(s: str | None):
    if not s:
        return None
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def short_id(ticket: dict) -> str:
    for a in ticket.get("aliases") or []:
        aid = a.get("id") or a.get("aliasId") or ""
        if RE_ALIAS.match(aid):
            return aid
    return ""


def attachments_summary(ticket: dict) -> str:
    att = ticket.get("attachments")
    items = att if isinstance(att, list) else (att.get("attachments") if isinstance(att, dict) else None)
    if not items:
        return "{'message': 'No attachments found.'}"
    counts: dict[str, int] = {}
    for a in items:
        name = (a.get("fileName") or a.get("name") or a.get("title") or "")
        ext = name.rsplit(".", 1)[-1].lower() if "." in name else (a.get("contentType") or "unknown")
        counts[ext] = counts.get(ext, 0) + 1
    return str({"total_attachments": len(items), "counts_by_type": counts})


def explode(tickets: list[dict]) -> list[dict]:
    rows = []
    sl = 0
    for t in tickets:
        desc = t.get("description") or ""
        tids = list(dict.fromkeys(RE_TID.findall(desc))) or ["TID not found"]
        m_station = RE_STATION.search(desc)
        m_zone = RE_ZONE.search(desc)
        sid = short_id(t)
        base = {
            "case_id": t.get("id"),
            "createdAt": parse_iso(t.get("createDate")),
            "lastUpdatedAt": parse_iso(t.get("lastUpdatedDate")),
            "station_code": m_station.group(1).upper() if m_station else "",
            "zone": m_zone.group(1) if m_zone else "",
            "attachments": attachments_summary(t),
            "status": (t.get("extensions") or {}).get("tt", {}).get("status") or t.get("status") or "",
            "IssueUrl": f"https://t.corp.amazon.com/issues/{sid}" if sid else "",
            "ShortId": sid,
        }
        for tid in tids:
            rows.append({"Sl": sl, **base, "impacted_tracking_ids": tid})
            sl += 1
    return rows


# --- write -------------------------------------------------------------------------------------
def write_outputs(rows: list[dict], stamp: str) -> tuple[Path, Path]:
    xlsx = OUT / f"Tech_Data_{stamp}.xlsx"
    csvp = OUT / f"Tech_Data_{stamp}.csv"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tech_Data"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r.get(c) for c in COLUMNS])
    for col in ("C", "D"):
        for cell in ws[col][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm:ss"
    for i, width in enumerate([6, 38, 20, 20, 12, 8, 20, 40, 18, 44, 14], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    wb.save(xlsx)
    with open(csvp, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        for r in rows:
            w.writerow({**r, "createdAt": r["createdAt"].isoformat(sep=" ") if r["createdAt"] else "",
                        "lastUpdatedAt": r["lastUpdatedAt"].isoformat(sep=" ") if r["lastUpdatedAt"] else ""})
    return xlsx, csvp


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--max-pages", type=int, default=None, help="debug: stop after N pages")
    ap.add_argument("--keep-old", action="store_true", help="do not delete previous Tech_Data_* files")
    args = ap.parse_args()

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    log(f"=== tickets_pull {stamp}")
    if not args.keep_old:
        for old in OUT.glob("Tech_Data_*.*"):
            old.unlink()
        log("deleted older Tech_Data_* files (delete-old-first rule)")

    mcp = open_mcp()
    try:
        tickets, total, pages = fetch_all(mcp, args.max_pages)
    finally:
        mcp.close()

    rows = explode(tickets)
    expected = sum(max(1, len(set(RE_TID.findall(t.get("description") or "")))) for t in tickets)
    n_no_tid = sum(1 for r in rows if r["impacted_tracking_ids"] == "TID not found")
    n_no_station = sum(1 for t in tickets if not RE_STATION.search(t.get("description") or ""))
    xlsx, csvp = write_outputs(rows, stamp)

    complete = (len(tickets) == total) or bool(args.max_pages)
    log(f"WROTE: {xlsx.name} data_rows={len(rows)} tickets={len(tickets)} total_at_fetch={total} "
        f"pages={pages} complete={complete}")
    log(f"checks: rows==expected_explosion {len(rows) == expected} ({len(rows)} vs {expected}); "
        f"tickets_without_TID={n_no_tid}; tickets_without_station_line={n_no_station}")
    if not complete:
        log("WARNING: collected != totalResults — do NOT treat this file as complete")
    print(json.dumps({"xlsx": str(xlsx), "csv": str(csvp), "rows": len(rows), "tickets": len(tickets),
                      "total_at_fetch": total, "complete": complete}))


if __name__ == "__main__":
    main()
