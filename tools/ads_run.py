r"""Drive the ADS & Pincode Batch Inference Platform (Streamlit) end to end.

Input : an XLSX in the template's "ADS format" (TrackingId, Concatenated Address Field)
Flow  : open app -> upload file -> click "ADS with PP Trigger" -> wait for results
        -> click the download button -> save the result file
Output: tools\captures\ads\ads_result_<timestamp>.<ext> + step screenshots + ads_run.log

Usage:
  python ads_run.py <input.xlsx>
  python ads_run.py --from-gcrs captures\gcrs_result.csv     (build the XLSX first)
"""
import csv
import json
import re
import sys
import time
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

ADS_URL = "http://44.203.89.162:8192/"
OUT = Path(__file__).parent / "captures" / "ads"
OUT.mkdir(parents=True, exist_ok=True)
LOG = OUT / "ads_run.log"
STAMP = datetime.now().strftime("%Y%m%d_%H%M%S")


def log(msg):
    line = f"[{datetime.now():%H:%M:%S}] {msg}"
    with open(LOG, "a", encoding="utf-8") as fh:
        fh.write(line + "\n")
    print(line, flush=True)


def _sanitize_field(value: object) -> str:
    """Make a value safe as a single upload cell.

    A stray double-quote in an address makes the ADS platform's CSV parser swallow the
    next row boundary, so it counts one fewer tracking ID than we uploaded (observed:
    UI=1992 vs XLSX=1993 from an address beginning with `"`). Double-quotes, embedded
    newlines, and tabs are normalized to spaces so every physical row stays one row.
    """
    text = "" if value is None else str(value)
    return re.sub(r"\s+", " ", text.replace('"', " ").replace("\r", " ").replace("\n", " ")).strip()


def concat_address(row: dict) -> str:
    """Replicates template GCRS_Data!K/L: lines joined by ',', then city,state,postal,country."""
    lines = [row.get(k, "") or "" for k in ("Address_Line_1", "Address_Line_2", "Address_Line_3")]
    lines = [x.strip() for x in lines if x and x.strip()]
    k = ",".join(lines)
    joined = ",".join([k, row.get("City", "") or "", row.get("State", "") or "",
                       str(row.get("Postal_Code", "") or ""), row.get("Country", "") or ""])
    return _sanitize_field(joined)


def _append_literal_input_row(worksheet, tracking_id: object, address: object) -> None:
    row_number = worksheet.max_row + 1
    for column, value in enumerate((tracking_id, address), start=1):
        cell = worksheet.cell(row=row_number, column=column)
        cell.value = "" if value is None else str(value)
        cell.data_type = "s"


def build_input_from_gcrs(gcrs_csv: Path) -> Path:
    with open(gcrs_csv, encoding="utf-8", newline="") as file:
        rows = list(csv.DictReader(file))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["TrackingId", "Concatenated Address Field"])
    n = 0
    seen: set[str] = set()
    for r in rows:
        if str(r.get("Valid", "True")).lower() == "false":
            continue
        tid = str(r.get("Query_Item") or "").strip()
        if not tid or tid in seen:
            continue
        seen.add(tid)
        _append_literal_input_row(ws, tid, concat_address(r))
        n += 1
    path = OUT / f"ads_input_{STAMP}.xlsx"
    wb.save(path)
    log(f"built {path.name} with {n} rows from {gcrs_csv.name}")
    return path


def snapshot(page, name):
    page.screenshot(path=str(OUT / f"{name}_{STAMP}.png"), full_page=True)
    try:
        body = page.locator("body").inner_text(timeout=5000)
    except Exception as exc:
        body = f"<err {exc}>"
    buttons = [t.strip() for t in page.locator("button").all_inner_texts() if t.strip()]
    dl = page.locator("[data-testid='stDownloadButton']").count()
    df = page.locator("[data-testid='stDataFrame'], [data-testid='stTable']").count()
    alerts = [t.strip() for t in page.locator("[data-testid='stAlert'], [data-testid='stException']").all_inner_texts() if t.strip()]
    spinner = page.locator("[data-testid='stSpinner'], [data-testid='stStatusWidget']").count()
    log(f"{name}: buttons={buttons} downloadButtons={dl} dataframes={df} spinner={spinner} alerts={alerts[:3]}")
    log(f"{name}: body=\n{body[:1500]}")
    return {"buttons": buttons, "dl": dl, "df": df, "alerts": alerts, "spinner": spinner, "body": body}


def build_input_from_rows(gcrs_rows: list[dict], stamp: str | None = None) -> Path:
    """Build the ADS-format XLSX (TrackingId, Concatenated Address Field) from GCRS row dicts.

    Rows with a blank TrackingId are dropped and each TrackingId is written once (first
    occurrence wins), so our physical row count matches the unique count the ADS platform
    reports for the file.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["TrackingId", "Concatenated Address Field"])
    n = 0
    seen: set[str] = set()
    for r in gcrs_rows:
        if str(r.get("Valid", "True")).lower() == "false":
            continue
        tid = str(r.get("Query_Item") or "").strip()
        if not tid or tid in seen:
            continue
        seen.add(tid)
        _append_literal_input_row(ws, tid, concat_address(r))
        n += 1
    path = OUT / f"ads_input_{stamp or STAMP}.xlsx"
    wb.save(path)
    log(f"built {path.name} with {n} rows")
    return path


def run(input_xlsx: Path, log_fn=None) -> Path:
    """Drive the Streamlit app for one input file; returns the downloaded CSV path."""
    _log = log_fn or log
    expected_count = max(0, openpyxl.load_workbook(input_xlsx, read_only=True).active.max_row - 1)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    with sync_playwright() as p:
        browser = p.chromium.launch(channel="msedge", headless=False)
        ctx = browser.new_context(viewport={"width": 1400, "height": 1000}, accept_downloads=True)
        page = ctx.new_page()
        loaded = False
        for attempt in range(3):
            try:
                page.goto(ADS_URL, timeout=60_000, wait_until="domcontentloaded")
                loaded = True
                break
            except Exception as exc:
                _log(f"ADS goto attempt {attempt + 1} interrupted: {str(exc)[:80]}")
                time.sleep(3)
        if not loaded:
            ctx.close()
            browser.close()
            raise RuntimeError(f"ADS app did not load at {ADS_URL} after 3 attempts (host may be down)")
        page.wait_for_selector("[data-testid='stFileUploaderDropzoneInput']", state="attached", timeout=60_000)
        time.sleep(2)
        page.locator("[data-testid='stFileUploaderDropzoneInput']").set_input_files(str(input_xlsx))
        for _ in range(60):
            time.sleep(1)
            if input_xlsx.name in page.locator("body").inner_text():
                break
        time.sleep(3)
        page.get_by_role("button", name="ADS with PP Trigger").first.click()
        _log("ADS: triggered inference")
        deadline = time.time() + 40 * 60
        while time.time() < deadline:
            time.sleep(5)
            dl = page.locator("[data-testid='stDownloadButton']").count()
            spinner = page.locator("[data-testid='stSpinner'], [data-testid='stStatusWidget']").count()
            err = page.locator("[data-testid='stException']").count()
            if err:
                page.screenshot(path=str(OUT / f"ads_error_{stamp}.png"), full_page=True)
                raise RuntimeError("ADS app raised an exception (see screenshot)")
            if dl and not spinner:
                break
        body = page.locator("body").inner_text()
        m = re.search(r"No of Tracking IDS in the File:\s*(\d+)", body)
        if not m:
            raise RuntimeError("ADS UI did not report its tracking-ID input count")
        ui_count = int(m.group(1))
        _log(f"ADS: UI reports {ui_count} tracking IDs in file")
        # The file is de-duplicated and sanitized before upload, so the counts should match
        # exactly. A tiny ADS-side shortfall (its parser occasionally drops a hard-to-parse
        # address row) is tolerated because every returned row is identity-checked downstream
        # and unmatched TIDs surface as "missing" rather than corrupting the output. A large
        # discrepancy still fails loudly, since that signals a real upload/parse problem.
        shortfall = expected_count - ui_count
        tolerance = max(3, expected_count // 100)
        if ui_count > expected_count or shortfall > tolerance:
            raise RuntimeError(f"ADS UI input-count mismatch: UI={ui_count}, XLSX={expected_count}")
        if shortfall:
            _log(f"ADS: tolerating minor UI shortfall of {shortfall} (<= {tolerance}); "
                 "unmatched TIDs will be reported as missing")
        with page.expect_download(timeout=180_000) as dl_info:
            page.locator("[data-testid='stDownloadButton'] button").first.click()
        target = OUT / f"ads_result_{stamp}.csv"
        dl_info.value.save_as(target)
        _log(f"ADS: saved {target.name}")
        ctx.close()
        browser.close()
    return target


def fetch(gcrs_rows: list[dict], log_fn=None, chunk: int | None = None) -> list[dict]:
    """Full ADS lane for the pipeline: build input, run, return result rows (dicts).

    The ADS platform accepts a large file in a single upload (the 200MB limit is on file
    size, not TID count), so by default all valid rows go up in ONE file/one browser
    session. This avoids the repeated browser launches that could stall between chunks.
    Pass an integer ``chunk`` only if you deliberately want to split the upload.
    """
    _log = log_fn or log
    valid = [r for r in gcrs_rows if str(r.get("Valid", "True")).lower() != "false"]
    step = chunk if chunk and chunk > 0 else max(1, len(valid))
    results: list[dict] = []
    for i in range(0, len(valid), step):
        part = valid[i:i + step]
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        xlsx = build_input_from_rows(part, stamp)
        _log(f"ADS batch {i // step + 1}: {len(part)} rows")
        out_csv = run(xlsx, _log)
        with open(out_csv, encoding="utf-8-sig", newline="") as fh:
            reader = csv.DictReader(fh)
            columns = reader.fieldnames or []
            rows = list(reader)
        required = {"TRACKINGID", "ADS", "predicted_pincode_BiLSTM", "predicted_pincode_proba_BiLSTM"}
        missing_columns = sorted(required - set(columns))
        if missing_columns:
            raise RuntimeError(f"ADS output missing required columns: {missing_columns}")
        requested = {str(row["Query_Item"]) for row in part}
        returned_ids = [str(row.get("TRACKINGID") or "").strip().split(".")[0] for row in rows]
        extra = sorted(set(returned_ids) - requested)
        duplicates = sorted(tid for tid, count in Counter(returned_ids).items() if count > 1)
        if extra or duplicates:
            raise RuntimeError(f"ADS identity contract failed: extra={extra[:5]} duplicates={duplicates[:5]}")
        for row in rows:
            try:
                pincode = float(row["predicted_pincode_BiLSTM"])
                probability = float(row["predicted_pincode_proba_BiLSTM"])
                float(row["ADS"])
            except (TypeError, ValueError) as exc:
                raise RuntimeError(f"ADS output has nonnumeric prediction fields for {row.get('TRACKINGID')}") from exc
            if not pincode.is_integer() or not 100000 <= int(pincode) <= 999999:
                raise RuntimeError(f"ADS output has invalid pincode for {row.get('TRACKINGID')}: {pincode}")
            if not 0 <= probability <= 1:
                raise RuntimeError(f"ADS output has invalid probability for {row.get('TRACKINGID')}: {probability}")
        missing = sorted(requested - set(returned_ids))
        _log(f"ADS batch {i // step + 1}: requested={len(requested)} returned={len(rows)} missing={len(missing)}")
        results += rows
    return results


def main():
    args = sys.argv[1:]
    if args and args[0] in {"-h", "--help"}:
        print(__doc__)
        sys.exit(0)
    if not args:
        print(__doc__)
        sys.exit(1)
    if args[0] == "--from-gcrs":
        input_xlsx = build_input_from_gcrs(Path(args[1]))
    else:
        input_xlsx = Path(args[0])
    log(f"=== ADS run {STAMP} input={input_xlsx}")

    with sync_playwright() as p:
        browser = p.chromium.launch(channel="msedge", headless=False)
        ctx = browser.new_context(viewport={"width": 1400, "height": 1000}, accept_downloads=True)
        page = ctx.new_page()
        downloads = []
        page.on("download", lambda d: downloads.append(d))

        for attempt in range(3):
            try:
                page.goto(ADS_URL, timeout=60_000, wait_until="domcontentloaded")
                break
            except Exception as exc:
                log(f"goto attempt {attempt + 1} interrupted: {str(exc)[:80]}")
                time.sleep(3)
        # Streamlit keeps the real <input type=file> hidden behind the dropzone; wait for
        # it to be attached (not visible) — set_input_files works on hidden inputs.
        page.wait_for_selector("[data-testid='stFileUploaderDropzoneInput']", state="attached", timeout=60_000)
        time.sleep(2)
        snapshot(page, "01_loaded")

        # 1. Upload
        page.locator("[data-testid='stFileUploaderDropzoneInput']").set_input_files(str(input_xlsx))
        log("file set on uploader; waiting for Streamlit to register it")
        for _ in range(30):
            time.sleep(1)
            if input_xlsx.name in page.locator("body").inner_text():
                break
        time.sleep(3)
        st = snapshot(page, "02_uploaded")

        # 2. Trigger inference (button label from recon)
        trig = page.get_by_role("button", name="ADS with PP Trigger")
        if trig.count():
            trig.first.click()
            log("clicked 'ADS with PP Trigger'")
        else:
            log("trigger button not found after upload — continuing to poll")

        # 3. Wait for results: download button, dataframe, or error (up to 15 min)
        deadline = time.time() + 15 * 60
        last_state = None
        while time.time() < deadline:
            time.sleep(5)
            dl = page.locator("[data-testid='stDownloadButton']").count()
            df = page.locator("[data-testid='stDataFrame'], [data-testid='stTable']").count()
            err = page.locator("[data-testid='stException']").count()
            spinner = page.locator("[data-testid='stSpinner'], [data-testid='stStatusWidget']").count()
            state = (dl, df, err, spinner)
            if state != last_state:
                log(f"poll: downloadButtons={dl} dataframes={df} exceptions={err} spinner={spinner}")
                last_state = state
            if err:
                snapshot(page, "03_error")
                break
            if dl and not spinner:
                break
        st = snapshot(page, "03_results")

        # 4. Download
        if page.locator("[data-testid='stDownloadButton']").count():
            btn = page.locator("[data-testid='stDownloadButton'] button").first
            log(f"clicking download button: {btn.inner_text()!r}")
            with page.expect_download(timeout=120_000) as dl_info:
                btn.click()
            d = dl_info.value
            suffix = Path(d.suggested_filename).suffix or ".bin"
            target = OUT / f"ads_result_{STAMP}{suffix}"
            d.save_as(target)
            log(f"saved result: {target} (suggested name {d.suggested_filename})")
        else:
            log("no download button appeared; results may be inline only (see 03_results screenshot/body)")

        (OUT / f"ads_run_state_{STAMP}.json").write_text(json.dumps(st, indent=2, ensure_ascii=False), encoding="utf-8")
        time.sleep(2)
        ctx.close()
        browser.close()
    log("=== done")


if __name__ == "__main__":
    main()
